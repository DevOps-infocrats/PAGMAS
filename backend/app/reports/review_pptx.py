"""Management Review deck (PowerPoint), built from analytics.review.management_review().

Layout follows the weekly governance review deck the client uses (20" × 11.25", blue title
band with the logo top-right, blue-headed tables with banded rows, week-on-week call-out box),
re-branded for Pratap. Every number comes from the review payload – nothing is typed in here.
Tables and charts are native PowerPoint objects, so the deck stays editable.
"""

from __future__ import annotations

import datetime as dt
import io
from functools import lru_cache
from typing import Any

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from ..config import PROJECT_DIR

LOGO = PROJECT_DIR / "frontend" / "public" / "brand" / "pratap-logo.png"
FONT = "Calibri"

# Pratap palette: brand blue dominates, logo green is the accent.
BRAND = "0B6BA8"
BRAND_DARK = "0A3A5F"
BRAND_LIGHT = "1E8FD0"
GREEN = "2BA84A"
INK = "16212E"
MUTED = "5B6675"
WHITE = "FFFFFF"
PANEL = "EEF4FA"
ROW_A = "EAF1F8"
ROW_B = "D9E5F1"
LINE = "BFD3E6"
GOOD = "16884A"
BAD = "C0392B"
BAND_COLORS = {
    ">98%": "0C8A0C",
    "95–98%": "5CB85C",
    "90–95%": "A8CF5F",
    "85–90%": "F6C445",
    "75–85%": "F2A65A",
    "50–75%": "EC835A",
    "30–50%": "E0583A",
    "0–30%": "C0392B",
    "0%": "7A1D1D",
}
CIRCLE_COLORS = [BRAND, GREEN, "E08A1E", "7B4FB3"]

SLIDE_W, SLIDE_H = 20.0, 11.25
MARGIN = 0.45
CONTENT_TOP = 1.9
CONTENT_BOTTOM = 10.55
TABLE_BOTTOM = 9.4  # tables stop above the call-out box


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


# ---------------------------------------------------------------- formatting


def n(v: Any) -> str:
    return "NA" if v is None else f"{v:,.0f}"


def pct(v: Any, digits: int = 2) -> str:
    return "NA" if v is None else f"{v:,.{digits}f}%"


def rkm(v: Any) -> str:
    return "NA" if v is None else f"{v:,.2f}"


def value(v: Any, unit: str) -> str:
    return pct(v) if unit == "percent" else rkm(v) if unit == "rkm" else n(v)


def delta_text(ch: dict[str, Any], unit: str = "count") -> str:
    d = ch.get("delta")
    if d is None:
        return "NA"
    if d == 0:
        return "■ 0" if unit != "percent" else "■ 0.00 pp"
    arrow = "▲" if d > 0 else "▼"
    return f"{arrow} {abs(d):,.2f} pp" if unit == "percent" else f"{arrow} {abs(d):,.0f}"


def trend_color(trend: str) -> str:
    return GOOD if trend == "better" else BAD if trend == "worse" else MUTED


def date_text(iso: str) -> str:
    return dt.date.fromisoformat(iso[:10]).strftime("%d %b %Y")


@lru_cache(maxsize=4)
def _logo_png(width_px: int) -> bytes | None:
    """Brand logo upscaled with Lanczos so it stays smooth at slide size."""
    if not LOGO.exists():
        return None
    im = Image.open(LOGO).convert("RGBA")
    h = round(im.height * width_px / im.width)
    buf = io.BytesIO()
    im.resize((width_px, h), Image.Resampling.LANCZOS).save(buf, format="PNG")
    return buf.getvalue()


# ---------------------------------------------------------------- primitives


class Deck:
    def __init__(self, review: dict[str, Any], generated_by: str):
        self.r = review
        self.generated_by = generated_by
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.prs.core_properties.title = "Management Review"
        self.prs.core_properties.author = "Pratap Network Analytics"
        self.prs.core_properties.subject = f"Week {review['week']['number']} management review"
        self.blank = self.prs.slide_layouts[6]
        self.page = 0

    # -- shapes
    def box(self, slide, x, y, w, h, fill=None, line=None, radius: float | None = None, line_w: float = 1.0, name: str | None = None):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        if radius:
            shape.adjustments[0] = radius
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(fill)
        else:
            shape.fill.background()
        if line:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Pt(line_w)
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        if name:
            shape.name = name
        shape.text_frame.text = ""
        return shape

    def text(
        self,
        slide,
        x,
        y,
        w,
        h,
        runs,
        size=14,
        color=INK,
        bold=False,
        align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP,
        margin=0.0,
        name: str | None = None,
        shape=None,
        line_spacing: float | None = None,
        space_after: float = 0,
    ):
        """runs: str | list of paragraphs; a paragraph is str or list of (text, {size,color,bold,italic})."""
        if shape is None:
            shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            shape.name = name
        tf = shape.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = anchor
        for side in ("left", "right", "top", "bottom"):
            setattr(tf, f"margin_{side}", Inches(margin))
        paras = runs if isinstance(runs, list) else [runs]
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            if line_spacing:
                p.line_spacing = line_spacing
            if space_after:
                p.space_after = Pt(space_after)
            for part in para if isinstance(para, list) else [(para, {})]:
                t, opt = part if isinstance(part, tuple) else (part, {})
                run = p.add_run()
                run.text = t
                f = run.font
                f.name = FONT
                f.size = Pt(opt.get("size", size))
                f.bold = opt.get("bold", bold)
                f.italic = opt.get("italic", False)
                f.color.rgb = rgb(opt.get("color", color))
        return shape

    def picture(self, slide, x, y, w):
        png = _logo_png(960)
        if png:
            slide.shapes.add_picture(io.BytesIO(png), Inches(x), Inches(y), width=Inches(w))

    # -- slide frame
    def new_slide(self, title: str, subtitle: str | None = None, notes: str | None = None):
        slide = self.prs.slides.add_slide(self.blank)
        self.page += 1
        self.box(slide, 0, 0, SLIDE_W, SLIDE_H, fill=WHITE, name="Background")
        self.box(slide, 0, 0, 17.85, 1.6, fill=BRAND, name="Title band")
        self.text(slide, MARGIN, 0.2, 17.1, 0.8, title, size=34, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, name="Title")
        if subtitle:
            self.text(slide, MARGIN, 0.98, 17.1, 0.45, subtitle, size=16, color="DCEBF7", name="Subtitle")
        self.picture(slide, 18.05, 0.12, 1.75)
        r = self.r
        footer = f"Pratap · Management Review · Week {r['week']['number']} ({date_text(r['week']['startDate'])} – {date_text(r['week']['endDate'])}) · Data as of {date_text(r['dataAsOf'])} · Source: {r['sourceFile']}"
        self.text(slide, MARGIN, 10.78, 16.5, 0.32, footer, size=10, color=MUTED, name="Footer")
        self.text(slide, 17.6, 10.78, 1.95, 0.32, f"{self.page}", size=10, color=MUTED, align=PP_ALIGN.RIGHT, name="Slide number")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return slide

    def callout(self, slide, x, y, w, h, label: str, items: list[dict[str, Any]], extra: str | None = None):
        """Week-on-week call-out box: "WK-40 → WK-41:  >98% GPs ▲ 3,340 → 3,423 | …"."""
        shape = self.box(slide, x, y, w, h, fill=PANEL, line=BRAND_LIGHT, radius=0.12, name="Call-out")
        parts: list[tuple[str, dict]] = [(f"{label}:  ", {"bold": True, "color": BRAND})]
        for i, it in enumerate(items):
            if i:
                parts.append(("   |   ", {"color": MUTED}))
            unit = it.get("unit", "count")
            arrow = "■" if not it["delta"] else ("▲" if it["delta"] > 0 else "▼")
            parts.append((f"{it['label']} ", {"color": INK}))
            parts.append((f"{arrow} ", {"color": trend_color(it["trend"]), "bold": True}))
            parts.append((f"{value(it['previous'], unit)} → {value(it['current'], unit)}", {"bold": True, "color": BRAND_DARK}))
        if extra:
            parts.append((f"   |   {extra}", {"color": INK}))
        self.text(slide, 0, 0, 0, 0, [parts], size=15, shape=shape, anchor=MSO_ANCHOR.MIDDLE, margin=0.2)
        return shape

    def section_label(self, slide, x, y, w, text: str, size=18):
        return self.text(slide, x, y, w, 0.45, text, size=size, bold=True, color=BRAND_DARK)

    # -- tables
    def table(
        self,
        slide,
        x,
        y,
        widths: list[float],
        header: list[list[str]],
        rows: list[list[Any]],
        row_h: float = 0.42,
        header_h: float = 0.5,
        size: int = 13,
        header_size: int | None = None,
        first_col_align=PP_ALIGN.LEFT,
        bold_rows: set[int] | None = None,
        total_rows: set[int] | None = None,
        cell_colors: dict[tuple[int, int], str] | None = None,
        merges: list[tuple[int, int, int, int]] | None = None,
        title: str | None = None,
        name: str | None = None,
        max_bottom: float = TABLE_BOTTOM,
        col_fills: dict[int, tuple[str, str, str]] | None = None,
        cell_fills: dict[tuple[int, int], tuple[str, str]] | None = None,
    ):
        """Blue header rows, banded body. rows are body rows; a body cell may be (text, colorHex).
        Body rows shrink (never below ~0.3") so the table ends above `max_bottom`.
        col_fills = {column: (header fill, body fill, body text colour)} to highlight whole columns;
        cell_fills = {(body row, column): (fill, text colour)} to flag single cells (wins over everything else)."""
        title_rows = 1 if title else 0
        head_total = header_h * (title_rows + len(header))
        if rows and y + head_total + row_h * len(rows) > max_bottom:
            row_h = max(0.3, (max_bottom - y - head_total) / len(rows))
            if row_h < 0.42:
                size = min(size, 12)
        n_rows = title_rows + len(header) + len(rows)
        n_cols = len(widths)
        heights = [header_h] * (title_rows + len(header)) + [row_h] * len(rows)
        gf = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(sum(widths)), Inches(sum(heights)))
        if name:
            gf.name = name
        tbl = gf.table
        tbl.first_row = False
        tbl.horz_banding = False
        for i, wdt in enumerate(widths):
            tbl.columns[i].width = Inches(wdt)
        for i, hgt in enumerate(heights):
            tbl.rows[i].height = Inches(hgt)
        hs = header_size or size

        def put(r, c, txt, fill, color, bold, align, sz):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = align
            run = p.add_run()
            run.text = "" if txt is None else str(txt)
            run.font.name = FONT
            run.font.size = Pt(sz)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)

        if title:
            put(0, 0, title, BRAND_LIGHT, WHITE, True, PP_ALIGN.CENTER, hs + 1)
            for c in range(1, n_cols):
                put(0, c, "", BRAND_LIGHT, WHITE, True, PP_ALIGN.CENTER, hs)
            tbl.cell(0, 0).merge(tbl.cell(0, n_cols - 1))
        for hr, hrow in enumerate(header):
            for c in range(n_cols):
                txt = hrow[c] if c < len(hrow) else ""
                head_fill = col_fills[c][0] if col_fills and c in col_fills else BRAND
                put(title_rows + hr, c, txt, head_fill, WHITE, True, PP_ALIGN.CENTER if c else first_col_align, hs)
        base = title_rows + len(header)
        for ri, row in enumerate(rows):
            total = bool(total_rows and ri in total_rows)
            fill = BRAND if total else (ROW_A if ri % 2 == 0 else ROW_B)
            for c in range(n_cols):
                cell = row[c] if c < len(row) else ""
                color = WHITE if total else INK
                if isinstance(cell, tuple):
                    cell, color = cell
                if cell_colors and (ri, c) in cell_colors and not total:
                    color = cell_colors[(ri, c)]
                cell_fill = fill
                if col_fills and c in col_fills:
                    cell_fill = col_fills[c][0] if total else col_fills[c][1]
                    if not total:
                        color = col_fills[c][2]
                if cell_fills and (ri, c) in cell_fills:
                    cell_fill, color = cell_fills[(ri, c)]
                bold = bool(total) or bool(bold_rows and ri in bold_rows) or c == 0 or bool(col_fills and c in col_fills)
                put(base + ri, c, cell, cell_fill, color, bold, first_col_align if c == 0 else PP_ALIGN.CENTER, size)
        for r0, c0, r1, c1 in merges or []:
            tbl.cell(r0, c0).merge(tbl.cell(r1, c1))
        return gf

    # -- charts
    def chart_style(self, chart, legend=True, size=12):
        chart.font.name = FONT
        chart.font.size = Pt(size)
        chart.font.color.rgb = rgb(MUTED)
        chart.has_legend = legend
        if legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(size)
        try:
            va = chart.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = rgb("E1E8EF")
            va.format.line.fill.background()
            va.tick_labels.font.size = Pt(size - 1)
            ca = chart.category_axis
            ca.format.line.color.rgb = rgb(LINE)
            ca.tick_labels.font.size = Pt(size - 1)
            ca.has_major_gridlines = False
        except (AttributeError, ValueError):
            pass

    def chart_title(self, chart, title: str):
        chart.has_title = True
        tf = chart.chart_title.text_frame
        tf.text = title
        f = tf.paragraphs[0].runs[0].font
        f.size = Pt(15)
        f.bold = True
        f.name = FONT
        f.color.rgb = rgb(BRAND_DARK)

    def band_chart(self, slide, x, y, w, h, title: str, periods: list[str], table: dict[str, Any]):
        cd = CategoryChartData()
        cd.categories = periods
        for row in reversed(table["bands"]):  # worst band at the bottom of the stack
            cd.add_series(row["band"], [v or 0 for v in row["values"]], number_format="#,##0")
        gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED_100, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        gf.name = "Band distribution chart"
        chart = gf.chart
        self.chart_style(chart, legend=True, size=11)
        self.chart_title(chart, title)
        plot = chart.plots[0]
        plot.gap_width = 55
        plot.overlap = 100
        for s in plot.series:
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = rgb(BAND_COLORS.get(s.name, BRAND))
            s.format.line.color.rgb = rgb(WHITE)
        chart.value_axis.tick_labels.number_format = "0%"
        chart.value_axis.tick_labels.number_format_is_linked = False
        return gf


# ---------------------------------------------------------------- slides


def _period_labels(r: dict[str, Any]) -> list[str]:
    return [p["label"] for p in r["periods"]]


def _cover(d: Deck):
    r = d.r
    slide = d.prs.slides.add_slide(d.blank)
    d.page += 1
    d.box(slide, 0, 0, SLIDE_W, SLIDE_H, fill=BRAND, name="Frame")
    d.box(slide, 0.5, 0.5, SLIDE_W - 1.0, SLIDE_H - 1.0, fill=WHITE, name="Card")
    d.picture(slide, 7.6, 0.95, 4.8)
    d.text(slide, 1.0, 5.55, 18.0, 1.0, "Management Review", size=54, bold=True, color=BRAND_DARK, align=PP_ALIGN.CENTER, name="Title")
    d.text(
        slide, 1.0, 6.55, 18.0, 0.6, "O&M Network Performance & Governance Review", size=26, color=BRAND, align=PP_ALIGN.CENTER, name="Subtitle"
    )
    w, c = r["week"], r["compareWeek"]
    d.text(
        slide,
        1.0,
        7.25,
        18.0,
        0.5,
        f"Week {w['number']} ({date_text(w['startDate'])} – {date_text(w['endDate'])}) compared with Week {c['number']}  ·  Data as of {date_text(r['dataAsOf'])}",
        size=17,
        color=MUTED,
        align=PP_ALIGN.CENTER,
        name="Period",
    )
    obj = d.box(slide, 1.5, 8.3, 17.0, 1.55, fill=BRAND, radius=0.1, name="Objective")
    d.text(
        slide,
        0,
        0,
        0,
        0,
        [[("Objective: ", {"bold": True, "size": 19}), ("A structured, data-driven weekly review of AMC conversion, targets and network availability across Rajasthan, Haryana and Assam – every figure is computed from the master data workbook.", {})]],
        size=17,
        color=WHITE,
        shape=obj,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.3,
    )
    d.text(slide, 1.5, 10.0, 17.0, 0.4, f"Source: {r['sourceFile']}  ·  Generated {dt.datetime.now():%d %b %Y %H:%M} by {d.generated_by}", size=11, color=MUTED, align=PP_ALIGN.CENTER)


def _index(d: Deck, sections: list[tuple[str, str]]):
    slide = d.new_slide("Key Index", "VISIBILITY  |  ACCOUNTABILITY  |  ACTION  |  CLOSURE")
    per_col = (len(sections) + 1) // 2
    for i, (title, desc) in enumerate(sections):
        col, row = divmod(i, per_col)
        x = 0.9 + col * 9.4
        y = 2.25 + row * 1.6
        d.box(slide, x, y, 8.9, 1.3, fill=PANEL, line=LINE, radius=0.15, name=f"Index {i + 1}")
        num = d.box(slide, x + 0.25, y + 0.25, 0.8, 0.8, fill=BRAND if i % 2 == 0 else GREEN, radius=0.5)
        d.text(slide, 0, 0, 0, 0, f"{i + 1:02d}", size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=num)
        d.text(slide, x + 1.3, y + 0.17, 7.4, 0.5, title, size=19, bold=True, color=BRAND_DARK)
        d.text(slide, x + 1.3, y + 0.67, 7.4, 0.45, desc, size=14, color=MUTED)


def _arrow(ch: dict[str, Any]) -> str:
    d = ch.get("delta")
    return "■" if not d else ("▲" if d > 0 else "▼")


def _change_text(ch: dict[str, Any], pp: float | None = None) -> str:
    """"▲ 151 (+3.9%) · +2.28 pp" – % change only when the base is non-zero."""
    d = ch.get("delta")
    if d is None:
        return "NA (no comparison value)"
    txt = f"{_arrow(ch)} {abs(d):,.0f}"
    if ch.get("pctChange") is not None:
        txt += f" ({'+' if ch['pctChange'] >= 0 else '−'}{abs(ch['pctChange']):.1f}%)"
    if pp is not None:
        txt += f" · {'+' if pp >= 0 else '−'}{abs(pp):.2f} pp"
    return txt


def _executive(d: Deck):
    """Page 1 – Executive Summary: availability KPIs, circle AMC conversions, AMC & target table, insights."""
    r = d.r
    sm = r["summary"]
    per = sm["periods"]
    wk, pw = per["week"], per["compareWeek"]
    slide = d.new_slide(
        "Executive Summary",
        f"{sm['kicker']}  |  {sm['subtitle']}",
        notes="All figures are calculated from the active workbook. " + " ".join(sm["definitions"]),
    )
    gen = dt.datetime.fromisoformat(r["generatedAt"]).astimezone().strftime("%d %b %Y %H:%M")
    d.text(
        slide,
        MARGIN,
        1.72,
        SLIDE_W - 2 * MARGIN,
        0.42,
        [[
            ("Current week ", {"color": MUTED}), (wk["title"], {"bold": True, "color": BRAND_DARK}),
            ("   vs   ", {"color": MUTED}), (pw["title"], {"bold": True, "color": BRAND_DARK}),
            ("     |     Month ", {"color": MUTED}), (per["month"]["range"], {"bold": True, "color": BRAND_DARK}),
            (" vs ", {"color": MUTED}), (per["previousMonth"]["range"], {"bold": True, "color": BRAND_DARK}),
            (f"     |     Data to {date_text(per['cutoff'])} · generated {gen}", {"color": MUTED}),
        ]],
        size=13,
        name="Reporting periods",
    )

    # Row 1 – three availability KPI cards: current week | previous week, then a plain change line
    cw = (SLIDE_W - 2 * MARGIN - 2 * 0.3) / 3

    def short(iso_a: str, iso_b: str) -> str:
        return f"{dt.date.fromisoformat(iso_a):%d %b} – {dt.date.fromisoformat(iso_b):%d %b}"

    for i, k in enumerate(sm["kpis"]):
        x, y = MARGIN + i * (cw + 0.3), 2.25
        d.box(slide, x, y, cw, 2.4, fill=WHITE, line=LINE, radius=0.06, name=f"KPI {k['key']}")
        d.text(slide, x + 0.2, y + 0.1, cw - 0.4, 0.4, k["label"], size=16, bold=True, color=BRAND_DARK)
        half = (cw - 0.55) / 2
        for j, (title, wkx, val, pc, den, strong) in enumerate(
            (("CURRENT WEEK", wk, k["current"], k["currentPct"], k["currentDenominator"], True), ("PREVIOUS WEEK", pw, k["previous"], k["previousPct"], k["previousDenominator"], False))
        ):
            bx = x + 0.2 + j * (half + 0.15)
            box = d.box(slide, bx, y + 0.55, half, 1.3, fill=PANEL if strong else "F4F6F9", line=BRAND_LIGHT if strong else LINE, radius=0.08)
            d.text(
                slide,
                0,
                0,
                0,
                0,
                [
                    [(title, {"size": 9, "bold": True, "color": MUTED})],
                    [(f"{wkx['label']} ", {"size": 12, "bold": True, "color": INK}), (f"({short(wkx['start'], wkx['end'])})", {"size": 10, "color": MUTED})],
                    [(n(val), {"size": 26, "bold": True, "color": BRAND if strong else INK})],
                    [(f"{pct(pc)} of {n(den)} reporting {k['unitOfCount']}", {"size": 9, "color": MUTED})],
                ],
                shape=box,
                margin=0.1,
                anchor=MSO_ANCHOR.MIDDLE,
            )
        verdict = {"better": "Improved", "worse": "Declined", "flat": "No change"}.get(k["trend"], "No comparison")
        change = "NA" if k["delta"] is None else f"{'+' if k['delta'] > 0 else '−' if k['delta'] < 0 else '±'}{abs(k['delta']):,.0f} {k['unitOfCount']}"
        if k.get("pctChange") is not None:
            change += f" ({'+' if k['pctChange'] >= 0 else '−'}{abs(k['pctChange']):.1f}%)"
        if k.get("ppChange") is not None:
            change += f" · {'+' if k['ppChange'] >= 0 else '−'}{abs(k['ppChange']):.2f} pp share"
        d.text(slide, x + 0.2, y + 1.95, cw - 0.4, 0.38, [[(f"{verdict}  ", {"bold": True, "color": trend_color(k["trend"])}), ("Change ", {"color": MUTED}), (change, {"bold": True, "color": INK})]], size=12, anchor=MSO_ANCHOR.MIDDLE)

    # Row 2 – Key highlights | AMC & month plan (left) + executive insights (right)
    ht = sm["highlightTable"]
    weeks = ht["weeks"]
    rows, colors = [], {}
    last_col = 1 + len(weeks)  # column of the reporting week
    for x in [*ht["rows"], ht["total"]]:
        rows.append([x["name"], n(x["currentAmc"]), *[n(v) for v in x["weeklyConversions"]], "—" if x["monthPlan"] is None else n(x["monthPlan"])])
    for ri in range(len(rows) - 1):
        colors[(ri, last_col)] = BRAND
    tw = 11.0
    k = len(weeks)
    widths = [1.75, 1.45] + [(tw - 1.75 - 1.45 - 1.3) / max(1, k)] * k + [1.3]
    group = ["Circle", f"Current AMC ({date_text(ht['cutoff'])[:6]})", "GPs achieved (per week)"] + [""] * (k - 1) + [ht["monthPlanLabel"]]
    names = ["", ""] + [w["label"] for w in weeks] + [""]
    d.table(
        slide,
        MARGIN,
        4.85,
        widths,
        [group, names],
        rows,
        row_h=0.62,
        header_h=0.5,
        size=15,
        header_size=11,
        total_rows={len(rows) - 1},
        cell_colors=colors,
        merges=[(1, 0, 2, 0), (1, 1, 2, 1), (1, 2, 1, 1 + k), (1, 2 + k, 2, 2 + k)],
        title=ht["title"],
        max_bottom=8.75,
        name="Key highlights",
    )
    note = ht["basis"]
    if sm["warnings"]:
        note += "  ⚠ " + " ".join(sm["warnings"])
    d.text(slide, MARGIN, 8.9, tw, 1.6, note, size=11, color=MUTED, name="Table notes")

    x0 = MARGIN + tw + 0.3
    box = d.box(slide, x0, 4.85, SLIDE_W - MARGIN - x0, 5.7, fill=PANEL, line=BRAND_LIGHT, radius=0.05, name="Executive insights")
    tags = {"observation": ("Fact", BRAND), "action": ("Action", GREEN), "data": ("Data", "B7791F")}
    paras = [[("Executive insights", {"bold": True, "size": 15, "color": BRAND_DARK})]]
    for it in sm["insights"]:
        tag, col = tags.get(it["kind"], ("Note", MUTED))
        paras.append([(f"{tag}: ", {"bold": True, "color": col}), (it["text"], {})])
    d.text(slide, 0, 0, 0, 0, paras, size=13, shape=box, margin=0.25, space_after=8)


def _targets_amc_cards(d: Deck, slide, y: float, h: float):
    """Rajasthan / Haryana / Assam – GPs converted to AMC (AMC + incremental, by AMC date), same figures as the web page."""
    cc = d.r["targetsPage"].get("circleCards")
    if not cc:
        return
    cw = (SLIDE_W - 2 * MARGIN - 2 * 0.3) / 3

    def signed(v: int) -> str:
        return f"{'+' if v > 0 else '−' if v < 0 else '±'}{abs(v):,}"

    for i, c in enumerate(cc["cards"]):
        x = MARGIN + i * (cw + 0.3)
        d.box(slide, x, y, cw, h, fill=WHITE, line=LINE, radius=0.07, name=f"Target AMC {c['code']}")
        chip = d.box(slide, x, y, cw, 0.42, fill=BRAND if i != 1 else GREEN, radius=0.07)
        d.text(slide, 0, 0, 0, 0, f"{c['name']} – GPs converted to AMC", size=14, bold=True, color=WHITE, shape=chip, anchor=MSO_ANCHOR.MIDDLE, margin=0.15)
        wk, mo = c["week"], c["month"]
        lines = [
            [("AMC GPs ", {"color": MUTED}), (n(c["amcNow"]), {"bold": True, "color": BRAND}), (f"  ({pct(c['amcPct'], 1)} of {n(c['revisedScope'])})", {"color": MUTED})]
            + ([("   ·   " + f"end of {c['baseline']['label']}: ", {"color": MUTED}), (n(c["baseline"]["value"]), {"bold": True, "color": INK}), (f" ({signed(c['baseline']['sinceChange'])})", {"bold": True, "color": GOOD})] if c.get("baseline") else []),
            [("Weekly  ", {"bold": True, "color": BRAND_DARK}), (f"{cc['weekLabel']} ", {"color": MUTED}), (n(wk["current"]), {"bold": True, "color": INK}), (f"  vs  {cc['compareLabel']} ", {"color": MUTED}), (n(wk["previous"]), {"bold": True, "color": INK}), (f"  ({signed(wk['delta'])})", {"bold": True, "color": trend_color(wk["trend"])})],
            [("Monthly  ", {"bold": True, "color": BRAND_DARK}), (f"{cc['monthLabel']} ", {"color": MUTED}), (n(mo["current"]), {"bold": True, "color": INK}), (f"  vs  {cc['previousMonthLabel']} ", {"color": MUTED}), (n(mo["previous"]), {"bold": True, "color": INK}), (f"  ({signed(mo['delta'])})", {"bold": True, "color": trend_color(mo["trend"])})],
            [("Per week  ", {"bold": True, "color": BRAND_DARK}), ("  ·  ".join(f"{w['label']} {w['value']:,}" for w in c["trend"]), {"color": INK})],
        ]
        d.text(slide, x + 0.2, y + 0.48, cw - 0.4, h - 0.52, lines, size=11, space_after=1)


def _targets(d: Deck):
    """Page 2 – Target & Achievement: circle AMC cards, monthly conversion plan, weekly closure targets."""
    r = d.r
    tp = r["targetsPage"]
    rm, wk = tp["roadmap"], tp["weekly"]
    slide = d.new_slide(
        "Target & Achievement",
        rm["subtitle"],
        notes=" ".join(tp["definitions"]) + (" Warnings: " + " ".join(tp["warnings"]) if tp["warnings"] else ""),
    )
    _targets_amc_cards(d, slide, 1.8, 1.45)

    # GP AMC Conversion & Liquidation Roadmap – open month plan, achieved so far (highlighted), next months' plans
    fut = rm["futureMonths"]
    ml = tp["monthLabel"]
    head = [f"{ml} Plan", rm["columns"].get("planAchieved", f"{ml} MTD Achieved"), *fut]

    def achieved_text(x: dict[str, Any]) -> str:
        v, pc = x.get("planAchieved"), x.get("planAchievedPct")
        return "—" if v is None else (f"{n(v)} ({pc:.1f}%)" if pc is not None else n(v))

    def plan_row(x: dict[str, Any], total: bool = False) -> list[Any]:
        cells = ["—" if x["mtdPlan"] is None else n(x["mtdPlan"]), achieved_text(x)]
        for i, v in enumerate(x["futurePlans"]):
            partial = total and x.get("futurePlansPartial", [False] * len(fut))[i]
            cells.append("—" if v is None else f"{n(v)}{'*' if partial else ''}")
        return [x["name"], *cells]

    rows = [plan_row(x) for x in rm["rows"]] + [plan_row(rm["total"], total=True)]
    lw = 6.9
    d.text(slide, MARGIN, 3.42, lw, 0.38, rm["title"], size=16, bold=True, color=BRAND_DARK, name="Roadmap title")
    first, plan_w, ach_w = 1.45, 1.0, 1.6
    rest = (lw - first - plan_w - ach_w) / max(1, len(fut))
    d.table(
        slide,
        MARGIN,
        3.85,
        [first, plan_w, ach_w] + [rest] * len(fut),
        [["Circle", *head]],
        rows,
        row_h=0.5,
        header_h=0.6,
        size=14,
        header_size=12,
        total_rows={len(rows) - 1},
        title="Monthly AMC conversion plan vs achieved",
        name="Liquidation roadmap",
        col_fills={2: (GOOD, "E3F4E8", GOOD)},
    )

    # Weekly closure targets: per week old (manual) target, new target (from the data, highlighted) and achieved
    weeks = wk["weeks"]

    def week_head(w: dict[str, Any]) -> str:
        span = f"{dt.date.fromisoformat(w['start']):%d %b}–{dt.date.fromisoformat(w['end']):%d %b}"
        return f"{w['label']} ({span}) · {'current week' if w.get('current') else w['status']}"

    head1 = ["Circle"] + sum([[week_head(w), "", ""] for w in weeks], [])
    head2 = [""] + ["Old", "New target", "Achieved"] * len(weeks)

    def ach_text(cell: dict[str, Any]) -> str:
        if cell["achieved"] is None:
            return "—"
        return f"{n(cell['achieved'])} ({cell['achievedPct']:.0f}%)" if cell.get("achievedPct") is not None else n(cell["achieved"])

    wrows = []
    flagged: dict[tuple[int, int], tuple[str, str]] = {}  # new target below the old target → red cell
    for ri, x in enumerate([*wk["rows"], wk["total"]]):
        line: list[Any] = [x["name"]]
        for wi, cell in enumerate(x["weeks"]):
            old = "—" if cell.get("oldTarget") is None else n(cell["oldTarget"])
            new = "—" if cell["target"] is None else n(cell["target"])
            change = cell.get("targetChange")
            if change is not None and change < 0:
                flagged[(ri, 2 + 3 * wi)] = (BAD, WHITE) if x["id"] == "all" else ("FBE3E0", BAD)
            line += [old if x["id"] == "all" else (old, MUTED), new, ach_text(cell)]
        wrows.append(line)
    x1 = MARGIN + lw + 0.35
    ww = SLIDE_W - MARGIN - x1
    circle_w = 1.3
    unit = (ww - circle_w) / max(1, len(weeks)) / 2.1  # old : new : achieved = 0.55 : 0.7 : 0.85
    widths = [circle_w] + [unit * 0.55, unit * 0.7, unit * 0.85] * len(weeks)
    merges = [(1, 0, 2, 0)] + [(1, 1 + 3 * i, 1, 3 + 3 * i) for i in range(len(weeks))]
    fills = {}
    for i in range(len(weeks)):
        fills[2 + 3 * i] = ("C98A00", "FFF4D6", "7A4A00")  # new target – amber
        fills[3 + 3 * i] = (GOOD, "E3F4E8", GOOD)  # achieved – green
    d.table(
        slide,
        x1,
        3.85,
        widths,
        [head1, head2],
        wrows,
        row_h=0.5,
        header_h=0.45,
        size=12 if len(weeks) <= 5 else 11,
        header_size=10,
        total_rows={len(wrows) - 1},
        merges=merges,
        title=wk["title"],
        max_bottom=7.3,
        name="Weekly closure targets",
        col_fills=fills,
        cell_fills=flagged,
    )

    note = d.box(slide, MARGIN, 7.55, SLIDE_W - 2 * MARGIN, 2.95, fill=PANEL, line=BRAND_LIGHT, radius=0.05, name="Target notes")
    paras = [[("How the figures are calculated", {"bold": True, "size": 13, "color": BRAND_DARK})]]
    paras += [[("•  ", {"color": BRAND, "bold": True}), (t, {})] for t in tp["definitions"]]
    paras += [[("⚠  ", {"color": "B7791F", "bold": True}), (w, {"color": "8A5A00"})] for w in tp["warnings"]]
    if any(rm["total"].get("futurePlansPartial", [])):
        paras.append([("* ", {"bold": True}), ("Overall is the sum of the circles that have a plan for that month.", {})])
    d.text(slide, 0, 0, 0, 0, paras, size=12, shape=note, margin=0.2, space_after=3)


def _operations(d: Deck):
    r = d.r
    ops = r["operations"]
    codes = [c["code"] for c in r["circles"]]
    slide = d.new_slide("Operations Snapshot", "Scope  |  HOTO Status  |  AMC Status  |  Gap – Scope vs AMC", notes="Gap = scope − AMC. Blocks under AMC = blocks with at least one AMC GP.")
    k = len(codes)
    header = [["Period", "Parameter", "Scope"] + [""] * (k - 1) + ["HOTO status"] + [""] * (k - 1) + ["AMC status"] + [""] * (k - 1) + ["Gap – scope vs AMC"] + [""] * (k - 1), ["", ""] + codes * 4]
    rows = []
    for p in ops:
        for pi, (label, key, fmt) in enumerate((("Total blocks", "blocks", n), ("Total GPs", "gps", n), ("Fibre RKM", "rkm", rkm))):
            row = [p["label"] if pi == 0 else "", label]
            for part in ("scope", "hoto", "amc", "gap"):
                row += [fmt(c[key][part]) for c in p["circles"]]
            rows.append(row)
    widths = [2.3, 1.6] + [round(15.1 / (4 * k), 3)] * (4 * k)
    merges = []
    for b in range(4):
        merges.append((0, 2 + b * k, 0, 1 + (b + 1) * k))
    merges += [(0, 0, 1, 0), (0, 1, 1, 1)]
    for i in range(len(ops)):
        merges.append((2 + i * 3, 0, 4 + i * 3, 0))
    d.table(slide, MARGIN, 1.95, widths, header, rows, row_h=0.56, header_h=0.5, size=14, merges=merges, name="Operations table")
    # overall chips
    sc = r["executive"]["scope"]
    y = 1.95 + 0.5 * 2 + 0.56 * len(rows) + 0.35
    for i, (label, key, fmt) in enumerate((("Total blocks", "blocks", n), ("Total GPs", "gps", n), ("Fibre RKM", "rkm", rkm))):
        x = MARGIN + i * 6.4
        d.box(slide, x, y, 6.1, 1.25, fill=PANEL, line=BRAND_LIGHT, radius=0.1)
        s = sc[key]
        amc_pct = s["amc"] / s["scope"] * 100 if s["scope"] else None
        hoto_pct = s["hoto"] / s["scope"] * 100 if s["scope"] else None
        d.text(slide, x + 0.2, y + 0.12, 5.7, 0.45, f"{label}  ·  Scope {fmt(s['scope'])}", size=16, bold=True, color=BRAND_DARK)
        d.text(slide, x + 0.2, y + 0.62, 5.7, 0.45, f"HOTO {pct(hoto_pct, 1)}   ·   AMC {pct(amc_pct, 1)}   ·   Gap {pct(None if amc_pct is None else 100 - amc_pct, 1)}", size=15, color=INK)


def _kpi_rows(section: dict[str, Any]) -> tuple[list[list[Any]], dict]:
    rows, colors = [], {}
    for k in section["kpis"]:
        label = k["label"] + (" *" if k["source"] != "computed" else "")
        rows.append([label] + [value(v, k["unit"]) for v in k["values"]])
    return rows, colors


def _critical_all(d: Deck):
    r = d.r
    sec = r["sections"]["all"]
    labels = _period_labels(r)
    slide = d.new_slide("Critical KPIs – All Circles", "Monthly & weekly performance  |  AMC GPs and blocks", notes="* One-Time-Up figures come from the circle sheets (external report).")
    rows, _ = _kpi_rows(sec)
    wl, wc = 3.75, min(0.95, 8.55 / max(1, len(labels)))  # 9 period columns at 0.95"; more columns share that width
    d.table(slide, MARGIN, 1.95, [wl] + [wc] * len(labels), [["Parameter"] + labels], rows, row_h=0.5, header_h=0.62, size=13, title="Monthly & Weekly Performance – All Circles", name="KPI table")
    # circle comparison for the selected week: rows = KPIs, column pairs = circles (value, change)
    wk = f"WK-{r['week']['number']}"
    pw = f"WK-{r['compareWeek']['number']}"
    keys = ["amcGps", "gpsAbove98", "gpsAtZero", "gpsBelowPenalty", "penaltyBlocks", "blocksAbove98", "availability", "naGps"]
    lab = {"amcGps": "AMC GPs", "gpsAbove98": ">98% GPs", "gpsAtZero": "0% GPs", "gpsBelowPenalty": "<90% GPs", "penaltyBlocks": "Blocks <90%", "blocksAbove98": ">98% blocks", "availability": "Avg avail.", "naGps": "NA GPs"}
    scopes = [*r["circles"], {"id": "all", "code": "ALL"}]
    changes = [{x["key"]: x for x in r["sections"][c["id"]]["weekChange"]} for c in scopes]
    crow, colors = [], {}
    for ri, k in enumerate(keys):
        row = [lab[k]]
        for si, ch in enumerate(changes):
            x = ch[k]
            row += [value(x["current"], x["unit"]), delta_text(x, x["unit"]).replace(" pp", "")]
            colors[(ri, 2 + si * 2)] = trend_color(x["trend"])
        crow.append(row)
    x0 = MARGIN + wl + wc * len(labels) + 0.25
    w = SLIDE_W - MARGIN - x0
    first = w * 0.16
    pair = (w - first) / len(scopes)
    d.table(
        slide,
        x0,
        1.95,
        [first] + [pair * 0.5, pair * 0.5] * len(scopes),
        [["KPI"] + sum([[c["code"], ""] for c in scopes], [])],
        crow,
        row_h=0.62,
        header_h=0.62,
        size=11,
        cell_colors=colors,
        merges=[(1, 1 + i * 2, 1, 2 + i * 2) for i in range(len(scopes))],
        title=f"{wk} vs {pw} (weekly columns; Δ = change, pp for %)",
        name="Circle comparison",
    )
    d.callout(slide, MARGIN, 9.6, SLIDE_W - 2 * MARGIN, 0.9, f"{pw} → {wk}", sec["insights"]["kpi"])


def _blocks_gps(d: Deck, c: dict[str, Any]):
    """Page 3 – Critical KPI – Blocks & GPs: monthly & weekly performance table for one circle."""
    ct = d.r["sections"][c["id"]]["critical"]
    slide = d.new_slide(
        f"Critical KPI – Blocks & GPs {c['name']}",
        "Monthly & weekly performance  |  GPs & blocks",
        notes="Months = workbook monthly AMC columns; MTD = month to date from daily data; weeks = source W-nn columns. "
        "Penalty-prone = block mean AMC availability below 90%. Previous-month MTD (e.g. Sep-26 MTD) = 1st to the same day of the previous month. "
        "One-Time-Up rows come from the circle sheet. "
        "Incremental GPs TD = incremental GPs whose AMC date falls in the column's period.",
    )
    labels = [col["label"] for col in ct["columns"]]
    rows = []
    for row in ct["rows"]:
        rows.append([row["label"] + (" *" if row["source"] != "computed" else ""), "" if row["scope"] is None else n(row["scope"]), *[n(v) for v in row["values"]]])
    mtd = next((i for i, col in enumerate(ct["columns"]) if col["kind"] == "mtd"), None)
    bold_cells = {}
    if mtd is not None:
        for ri in range(len(rows)):
            bold_cells[(ri, 2 + mtd)] = BRAND_DARK
    w = SLIDE_W - 2 * MARGIN
    first, scope = 5.2, 1.6
    d.table(
        slide,
        MARGIN,
        1.95,
        [first, scope] + [(w - first - scope) / len(labels)] * len(labels),
        [["Parameters", "Scope / Target", *labels]],
        rows,
        row_h=0.78,
        header_h=0.7,
        size=16,
        header_size=15,
        cell_colors=bold_cells,
        title=ct["title"],
        max_bottom=9.2,
        name="Monthly & weekly performance",
    )
    ins = ct["insight"]
    d.callout(slide, MARGIN, 9.5, w, 0.95, f"{ins['from']} → {ins['to']}", ins["items"], extra="* One-Time-Up: circle sheet")


def _band_rows(table: dict[str, Any], kind: str) -> tuple[list[list[Any]], set[int]]:
    rows = []
    if kind == "gp":
        rows.append(["Scope (AMC GPs)"] + [n(v) for v in table["scope"]])
    for b in table["bands"]:
        rows.append([b["band"]] + [n(v) for v in b["values"]])
    rows.append(["Total reported"] + [n(v) for v in table["total"]])
    rows.append(["NA (not reported)"] + [n(v) for v in table["na"]])
    rows.append([f"Below 90% ({'GPs' if kind == 'gp' else 'blocks'})"] + [n(v) for v in table["belowPenalty"]])
    rows.append(["Average availability"] + [pct(v) for v in table["average"]])
    bold = {0, len(rows) - 4} if kind == "gp" else {len(rows) - 4}
    return rows, bold


def _block_uptime(d: Deck, c: dict[str, Any]):
    """Page 4 – same content as the dashboard: the block band table for one circle + week call-out."""
    r = d.r
    sec = r["sections"][c["id"]]
    labels = _period_labels(r)
    slide = d.new_slide(
        f"Block Network Availability & Uptime Performance – {c['name']}",
        f"Full availability distribution (AMC blocks) — focus: blocks below {r['thresholds']['penalty']:g}%",
        notes="Block value = mean availability of its AMC GPs. NA = HOTO-done block without any reported AMC GP value.",
    )
    rows, bold = _band_rows(sec["blockUptime"], "block")
    w = SLIDE_W - 2 * MARGIN
    wl = 3.6
    d.table(
        slide,
        MARGIN,
        1.95,
        [wl] + [(w - wl) / len(labels)] * len(labels),
        [["Block uptime"] + labels],
        rows,
        row_h=0.5,
        header_h=0.62,
        size=14,
        bold_rows=bold,
        title=f"Block network availability & uptime – {c['name']}",
        name="Block band table",
    )
    d.callout(slide, MARGIN, 9.6, w, 0.9, f"WK-{r['compareWeek']['number']} → WK-{r['week']['number']}", sec["insights"]["block"])


def _gp_uptime(d: Deck, c: dict[str, Any]):
    """Page 5 – same content as the dashboard: GP band table (months, TD, weeks oldest → newest) + call-out."""
    r = d.r
    sec = r["sections"][c["id"]]
    slide = d.new_slide(
        f"GP Network Availability & Uptime Performance – {c['name']}",
        "GP performance – AMC GPs",
        notes="A GP counts in a column only when it is under AMC by the column end. NA = under AMC but no value in the source (never 0%).",
    )
    periods = r["periods"]
    head = [i for i, p in enumerate(periods) if p["kind"] != "week"]
    weeks = sorted((i for i, p in enumerate(periods) if p["kind"] == "week"), key=lambda i: periods[i]["endDate"])
    order = head + weeks
    labels = [f"WK-{periods[i]['key'].removeprefix('W-')}" if periods[i]["kind"] == "week" else periods[i]["label"] for i in order]
    rows, bold = _band_rows(sec["gpUptime"], "gp")
    rows = [[row[0]] + [row[1 + i] for i in order] for row in rows]
    w = SLIDE_W - 2 * MARGIN
    wl = 3.6
    d.table(
        slide,
        MARGIN,
        1.95,
        [wl] + [(w - wl) / len(labels)] * len(labels),
        [["Bucket"] + labels],
        rows,
        row_h=0.47,
        header_h=0.62,
        size=14,
        bold_rows=bold,
        title=f"GP network availability & uptime – {c['name']}",
        name="GP band table",
    )
    d.callout(slide, MARGIN, 9.6, w, 0.9, f"WK-{r['compareWeek']['number']} → WK-{r['week']['number']}", sec["insights"]["gp"])


def _data_notes(d: Deck):
    r = d.r
    dn = r["dataNotes"]
    slide = d.new_slide("Data Notes, Definitions & Reconciliation", "How every number in this deck is calculated and checked")
    d.box(slide, MARGIN, 1.95, 9.4, 4.15, fill=PANEL, line=LINE, radius=0.05)
    d.section_label(slide, MARGIN + 0.25, 2.05, 9, "Definitions")
    d.text(slide, MARGIN + 0.25, 2.55, 8.9, 3.5, [[("•  ", {"color": BRAND, "bold": True}), (t, {})] for t in dn["definitions"]], size=13, space_after=5)
    dq = [[q["label"], n(q["count"])] for q in dn["dataQuality"]][:9]
    if dq:
        d.table(slide, MARGIN, 6.35, [7.6, 1.8], [["Data-quality check (active file)", "Count"]], dq, row_h=0.42, header_h=0.5, size=12, max_bottom=10.55, name="Data quality")
    rec = dn["reconciliation"]
    cnt = rec["counts"]
    x0 = 10.15
    w = SLIDE_W - MARGIN - x0
    d.section_label(slide, x0, 1.95, w, "Reconciliation with the O&M KEY NUMBERS")
    d.text(
        slide,
        x0,
        2.45,
        w,
        1.0,
        f"{cnt.get('MATCH', 0)} figures match exactly · {cnt.get('EXPLAINED', 0)} differ with a documented reason · {cnt.get('MISMATCH', 0)} unexplained · {cnt.get('NO_REFERENCE', 0)} without a source value."
        + (f" The O&M month-to-date column covers data to {date_text(rec['sourceMtdEnd'])}; this review uses daily data to {date_text(r['dataAsOf'])}." if rec.get("sourceMtdEnd") else ""),
        size=13,
        color=INK,
    )
    status_color = {"MATCH": GOOD, "EXPLAINED": "B7791F", "MISMATCH": BAD, "NO_REFERENCE": MUTED}
    order = {"MISMATCH": 0, "EXPLAINED": 1, "NO_REFERENCE": 2, "MATCH": 3}
    recs = sorted(rec["rows"], key=lambda x: (order.get(x["status"], 9), x["state"]))[:14]
    rrows, colors = [], {}
    for i, x in enumerate(recs):
        unit = "rkm" if "rkm" in x["metric"].lower() or "fibre" in x["metric"].lower() else "count"
        rrows.append([x["state"], x["metric"], value(x["reported"], unit), value(x["computed"], unit), x["status"].replace("_", " ").title()])
        colors[(i, 4)] = status_color.get(x["status"], INK)
    if rrows:
        d.table(slide, x0, 3.55, [w * 0.16, w * 0.4, w * 0.14, w * 0.14, w * 0.16], [["Circle", "O&M metric", "O&M", "Review", "Status"]], rrows, row_h=0.45, header_h=0.5, size=12, cell_colors=colors, max_bottom=10.55, name="Reconciliation")


def _thanks(d: Deck):
    slide = d.prs.slides.add_slide(d.blank)
    d.page += 1
    d.box(slide, 0, 0, SLIDE_W, SLIDE_H, fill=BRAND)
    d.box(slide, 0.5, 0.5, SLIDE_W - 1.0, SLIDE_H - 1.0, fill=WHITE)
    d.picture(slide, 7.9, 2.0, 4.2)
    d.text(slide, 1, 6.6, 18, 1.0, "Thank You", size=50, bold=True, color=BRAND_DARK, align=PP_ALIGN.CENTER)
    d.text(slide, 1, 7.6, 18, 0.5, "Governed by data · Driven by accountability", size=20, color=BRAND, align=PP_ALIGN.CENTER)


# Deck sections the user can pick before downloading (order = slide order).
SECTIONS: dict[str, tuple[str, str]] = {
    # The Management Review report (pages 1–5)
    "summary": ("Executive Summary", "Availability KPIs, AMC & monthly target, insights"),
    "targets": ("Target & Achievement", "Circle AMC conversions, monthly conversion plan, weekly closure targets"),
    "blocks-gps": ("Critical KPI – Blocks & GPs", "Monthly & weekly performance table – {names}"),
    "block-network": ("Block Network Availability & Uptime Performance", "Block availability band table – {names}"),
    "gp-network": ("GP Network Availability & Uptime Performance", "GP availability band table – {names}"),
    # Optional extra pages
    "kpis": ("Critical KPIs – All Circles", "Monthly & weekly performance, circle comparison"),
    "operations": ("Operations Snapshot", "Scope, HOTO, AMC and gap for blocks, GPs and fibre"),
    "notes": ("Data Notes & Reconciliation", "Definitions, source checks and data quality"),
}
PER_CIRCLE = {"blocks-gps": _blocks_gps, "block-network": _block_uptime, "gp-network": _gp_uptime}
FIXED = {"summary": _executive, "targets": _targets, "operations": _operations, "kpis": _critical_all, "notes": _data_notes}


def render_review_deck(review: dict[str, Any], generated_by: str, sections: list[str] | None = None, circle_ids: list[str] | None = None) -> bytes:
    """Cover, index, the chosen sections (default: all) for the chosen circles (default: all), thank-you."""
    chosen = [s for s in SECTIONS if sections is None or s in sections]
    circles = [c for c in review["circles"] if circle_ids is None or c["id"] in circle_ids]
    if not chosen:
        raise ValueError("Choose at least one section")
    if not circles and any(s in PER_CIRCLE for s in chosen):
        raise ValueError("Choose at least one circle for the per-circle sections")
    d = Deck(review, generated_by)
    names = ", ".join(c["name"] for c in circles)
    _cover(d)
    _index(d, [(SECTIONS[s][0], SECTIONS[s][1].format(names=names)) for s in chosen])
    for s in chosen:
        if s in PER_CIRCLE:
            for c in circles:
                PER_CIRCLE[s](d, c)
        else:
            FIXED[s](d)
    _thanks(d)
    buf = io.BytesIO()
    d.prs.save(buf)
    return buf.getvalue()
