"""Management Review (web page + PPT) – figures checked against the workbook's own sheets,
plus the .env-driven admin account and JWT key rotation."""

import io

import jwt
import pytest
from pptx import Presentation

from app.analytics.review import management_review
from app.config import get_settings
from app.db import SessionLocal
from app.main import _ensure_env_admin
from app.models import User
from app.security.core import create_session_token, decode_session_token, verify_password

from .conftest import login


@pytest.fixture(scope="module")
def plan_file():
    """Tests run on a temporary data folder – give it the project's management plan."""
    import shutil
    from pathlib import Path

    src = Path(__file__).resolve().parents[2] / "data" / "review_plan.json"
    if not src.exists():
        pytest.skip("data/review_plan.json not available")
    dst = get_settings().data_dir / "review_plan.json"
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(src, dst)
    return dst


@pytest.fixture(scope="module")
def review(imported, plan_file):
    return management_review(imported, "W-41", "W-40", [])


def _row(section, key):
    return next(k for k in section["kpis"] if k["key"] == key)


def test_periods_follow_the_om_layout(review):
    assert [p["label"] for p in review["periods"]] == ["Jun-26", "Jul-26", "Aug-26", "Sep-26", "Sep-26 TD", "Oct-26 TD", "Till W-41", "WK-40", "WK-39", "WK-38"]
    assert [c["code"] for c in review["circles"]] == ["RJ", "HR", "AS"]  # by size, as in the deck


def test_rajasthan_kpis_match_om_sheet(review):
    rj = review["sections"]["rajasthan"]
    # O&M "AMC GPs Count" row (months, Till W-41, WK-40, WK-39, WK-38); the two TD columns are computed to the data date
    amc = _row(rj, "amcGps")["values"]
    assert amc[:4] == [3263, 3818, 4353, 4710] and amc[6:] == [4710, 4710, 4626, 4590]
    # O&M "GP Uptime Performance" >98% row – every column except TD (source TD stops at 6 Oct)
    above = _row(rj, "gpsAbove98")["values"]
    assert above[:4] == [1782, 2081, 2422, 2466] and above[6:] == [3423, 3340, 3249, 3145]


def test_gp_band_table_matches_om_till_week(review):
    rj = review["sections"]["rajasthan"]["gpUptime"]
    till = {b["band"]: b["values"][6] for b in rj["bands"]}
    assert till == {">98%": 3423, "95–98%": 309, "90–95%": 275, "85–90%": 169, "75–85%": 186, "50–75%": 204, "30–50%": 99, "0–30%": 37, "0%": 7}
    assert rj["total"][6] == 4709 and rj["na"][6] == 1


def test_block_bands_include_amc_blocks_without_hoto_date(review):
    """6 Assam blocks report AMC availability but have no HOTO date – they must be counted (O&M: 41)."""
    assert review["sections"]["assam"]["blockUptime"]["total"][6] == 41
    rj = {b["band"]: b["values"][6] for b in review["sections"]["rajasthan"]["blockUptime"]["bands"]}
    assert (rj[">98%"], rj["95–98%"], rj["90–95%"]) == (99, 79, 69)  # O&M Block Uptime, Till W-41


def test_one_time_up_comes_from_circle_sheet(review):
    rj = _row(review["sections"]["rajasthan"], "oneTimeUpGps")
    assert rj["source"] == "circle sheet"
    assert rj["values"][:4] == [6142, 5976, 5703, 5635] and rj["values"][6] == 5635
    assert rj["values"][4] is None  # Sep-26 TD: not in the circle sheet, never estimated


def test_targets_from_om_key_numbers(review):
    rows = {r["code"]: r for r in review["targets"]["rows"]}
    assert (rows["RJ"]["target"], rows["HR"]["target"], rows["AS"]["target"]) == (6359, 2729, 230)
    assert rows["RJ"]["targetGap"] == 6359 - 4710
    assert rows["ALL"]["target"] == 6359 + 2729 + 230
    assert rows["RJ"]["rfpTarget"] == 8746 * 0.75  # O&M "75% of Existing GPs (RFP)"
    assert review["targets"]["label"] == "Operational GP target (Oct-26)"


def test_overall_equals_sum_of_circles(review):
    sec = review["sections"]
    for key in ("amcGps", "gpsAbove98", "gpsAtZero", "penaltyBlocks", "naGps"):
        total = _row(sec["all"], key)["values"]
        parts = [_row(sec[c["id"]], key)["values"] for c in review["circles"]]
        assert total == [sum(p[i] for p in parts) for i in range(len(total))], key


def test_review_api_and_pptx(imported, client):
    login(client)
    r = client.get("/api/v1/reviews/management", params={"week": "W-41", "compare": "W-40"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["week"]["id"] == "W-41" and set(body["sections"]) == {"rajasthan", "haryana", "assam", "all"}
    assert client.get("/api/v1/reviews/management", params={"week": "W-40", "compare": "W-41"}).status_code == 422

    p = client.get("/api/v1/reviews/management.pptx", params={"week": "W-41", "compare": "W-40"})
    assert p.status_code == 200
    assert p.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml")
    assert 'filename="Pratap_Management_Review_W-41_vs_W-40.pptx"' in p.headers["content-disposition"]
    deck = Presentation(io.BytesIO(p.content))
    texts = ["\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame) for slide in deck.slides]
    assert len(deck.slides) == 17
    assert "Management Review" in texts[0]
    assert any(s.shape_type == 13 for s in deck.slides[0].shapes)  # logo picture on the cover
    for title in ("Executive Summary", "Target & Achievement", "Critical KPI – Blocks & GPs Rajasthan", "Block Network Availability & Uptime Performance", "GP Network Availability & Uptime Performance"):
        assert any(title in t for t in texts), title


def test_review_requires_session(client):
    assert client.get("/api/v1/reviews/management").status_code == 401
    assert client.get("/api/v1/reviews/management.pptx").status_code == 401


# ---------------------------------------------------------------- .env admin + JWT


@pytest.fixture
def settings_override():
    s = get_settings()
    saved = {k: getattr(s, k) for k in ("admin_email", "admin_password", "admin_password_hash", "admin_sync_password", "secret_key", "secret_key_previous")}
    yield s
    for k, v in saved.items():
        object.__setattr__(s, k, v)


def test_env_admin_created_and_synced(settings_override):
    s = settings_override
    object.__setattr__(s, "admin_email", "envadmin@pratap-test.in")
    object.__setattr__(s, "admin_password", "First!Env-Pass-123")
    with SessionLocal() as db:
        _ensure_env_admin(db)
        u = db.query(User).filter_by(email="envadmin@pratap-test.in").one()
        assert u.role == "admin" and verify_password("First!Env-Pass-123", u.password_hash)
        version = u.token_version

        object.__setattr__(s, "admin_password", "Second!Env-Pass-456")
        _ensure_env_admin(db)  # sync disabled -> unchanged
        db.refresh(u)
        assert verify_password("First!Env-Pass-123", u.password_hash)

        object.__setattr__(s, "admin_sync_password", True)
        _ensure_env_admin(db)
        db.refresh(u)
        assert verify_password("Second!Env-Pass-456", u.password_hash) and u.token_version == version + 1

        object.__setattr__(s, "admin_password", "weak")
        _ensure_env_admin(db)  # weak password rejected, account untouched
        db.refresh(u)
        assert verify_password("Second!Env-Pass-456", u.password_hash)


def test_jwt_key_rotation(settings_override):
    s = settings_override
    token, _ = create_session_token(1, "admin", 0)
    old_key = s.secret_key
    object.__setattr__(s, "secret_key", "n3w-" + old_key[::-1])
    object.__setattr__(s, "secret_key_previous", old_key)
    assert decode_session_token(token)["sub"] == "1"  # still valid during rotation
    object.__setattr__(s, "secret_key_previous", None)
    assert decode_session_token(token) is None  # old key retired


def test_jwt_wrong_audience_or_issuer_rejected():
    s = get_settings()
    token, _ = create_session_token(1, "admin", 0)
    claims = jwt.decode(token, options={"verify_signature": False})
    for field in ("aud", "iss"):
        bad = dict(claims, **{field: "someone-else"})
        assert decode_session_token(jwt.encode(bad, s.secret_key, algorithm=s.jwt_algorithm)) is None


def test_pptx_contains_only_the_chosen_sections_and_circles(imported, client):
    login(client)
    p = client.get("/api/v1/reviews/management.pptx", params={"week": "W-41", "compare": "W-40", "sections": "summary,gp-network", "circles": "haryana"})
    assert p.status_code == 200, p.text
    deck = Presentation(io.BytesIO(p.content))
    titles = [next((s.text_frame.text for s in slide.shapes if s.name == "Title"), "") for slide in deck.slides]
    # cover, index, executive summary, GP network – Haryana only, thank-you
    assert len(deck.slides) == 5
    assert titles[2] == "Executive Summary"
    assert titles[3] == "GP Network Availability & Uptime Performance – Haryana"
    for bad in ({"sections": "summary,hack"}, {"circles": "goa"}, {"sections": "gp-network", "circles": ""}):
        assert client.get("/api/v1/reviews/management.pptx", params={"week": "W-41", "compare": "W-40", **bad}).status_code == 422, bad


# ---------------------------------------------------------------- Page 1: Executive Summary


def test_reporting_periods_are_derived_from_the_file(imported):
    import datetime as dt

    from app.analytics.review_periods import latest_completed_week, month_window, reporting_periods

    done = latest_completed_week(imported)
    assert done.end <= imported.data_as_of and done.code == "W-41"
    p = reporting_periods(imported, done, imported.week("W-40"))
    assert p.cutoff == imported.data_as_of and p.month == (dt.date(2026, 10, 1), dt.date(2026, 10, 8))
    assert p.previous_month == (dt.date(2026, 9, 1), dt.date(2026, 9, 8)) and not p.warnings
    assert month_window(dt.date(2026, 10, 31), 1) == (dt.date(2026, 9, 1), dt.date(2026, 9, 30))  # shorter month clamps
    assert reporting_periods(imported, done, imported.week("W-38")).warnings  # not the immediate predecessor


def test_summary_kpis_recomputed_from_gp_rows(review, imported):
    """Availability KPI cards = independent recount of the weekly W-nn values."""
    sm = review["summary"]
    k = {x["key"]: x for x in sm["kpis"]}
    for code, side in (("W-41", "current"), ("W-40", "previous")):
        wf = imported.week_frame(code)
        v = wf["weekly"].where(wf["amc"])
        blocks = v.groupby(imported.gps["block_id"]).mean().dropna()
        assert k["gpsAbove98"][side] == int((v > 98).sum())
        assert k["gpsAtZero"][side] == int((v == 0).sum())  # NA never counted as 0
        assert k["blocksAbove98"][side] == int((blocks > 98).sum())
    assert k["gpsAbove98"]["currentPct"] == round(k["gpsAbove98"]["current"] / k["gpsAbove98"]["currentDenominator"] * 100, 2)
    assert k["gpsAtZero"]["trend"] == ("worse" if k["gpsAtZero"]["delta"] > 0 else "better" if k["gpsAtZero"]["delta"] < 0 else "flat")


def test_summary_conversions_split_amc_and_incremental(review, imported):
    import pandas as pd

    sm = review["summary"]
    amc = pd.to_datetime(imported.gps["amc_date"], errors="coerce")
    for card in sm["conversions"]:
        sel = imported.gps["circle"].str.title() == card["name"]
        inside = sel & (amc >= "2026-09-25") & (amc <= "2026-10-01")  # WK-40
        assert card["week"]["amc"]["previous"] == int((inside & (imported.gps["amc_category"] == "AMC")).sum())
        assert card["week"]["incremental"]["previous"] == int((inside & (imported.gps["amc_category"] == "INCREMENTAL")).sum())
    rj = next(c for c in sm["conversions"] if c["code"] == "RJ")
    assert rj["week"]["incremental"]["previous"] == 84 and rj["month"]["incremental"]["previous"] == 167
    assert rj["week"]["amc"]["pctChange"] is None  # no % change from a zero base
    assert all(c["ok"] for c in sm["checks"]) and sm["valid"]


def test_summary_target_table(review):
    ht = review["summary"]["highlightTable"]
    assert ht["targetLabel"] == "Oct-26 target" and ht["targetMonthMatches"]
    rows = {r["code"]: r for r in ht["rows"]}
    assert (rows["RJ"]["target"], rows["HR"]["target"], rows["AS"]["target"]) == (6359, 2729, 230)
    assert rows["RJ"]["amcGpsAsOfCutoff"] == 4710 and rows["RJ"]["targetProgressPct"] == round(4710 / 6359 * 100, 2)
    assert 3 <= len(review["summary"]["insights"]) <= 5


def test_review_defaults_to_latest_completed_week(imported, client):
    login(client)
    body = client.get("/api/v1/reviews/management").json()
    assert body["week"]["id"] == "W-41" and body["compareWeek"]["id"] == "W-40"
    assert body["summary"]["periods"]["week"]["complete"] is True


# ---------------------------------------------------------------- Page 2: Target & Achievement


def test_roadmap_formulas(review, imported):
    import pandas as pd

    tp = review["targetsPage"]
    rows = {r["code"]: r for r in tp["roadmap"]["rows"]}
    amc = pd.to_datetime(imported.gps["amc_date"], errors="coerce")
    for code, state, proj in (("RJ", "RAJASTHAN", 6359), ("HR", "HARYANA", 2729), ("AS", "ASSAM", 230)):
        r, sel = rows[code], imported.gps["circle"] == state
        assert r["revisedScope"] == int(imported.gps.loc[sel, "revised"].sum())
        assert r["amcTillPrevMonth"] == int((sel & (amc <= "2026-09-30")).sum())
        assert r["projection"] == proj  # O&M "Target for Operational GP (Oct-26)"
        assert r["planByMonth"] == r["amcTillPrevMonth"] + r["mtdPlan"]
        assert r["gap"] == max(0, proj - r["cumulativeMtd"])
        assert r["pendingAfterMonth"] == r["revisedScope"] - proj
    assert (rows["RJ"]["mtdPlan"], rows["HR"]["mtdPlan"], rows["AS"]["mtdPlan"]) == (1063, 652, 60)  # management plan
    assert (rows["HR"]["readyForIncremental"], rows["AS"]["readyForIncremental"]) == (310, 58)  # letter submitted, not yet AMC
    total = tp["roadmap"]["total"]
    assert tp["roadmap"]["futureMonths"] == ["Nov-26", "Dec-26", "Jan-27"]
    assert total["futurePlans"] == [2505, 2503, 1061] and total["futurePlansPartial"] == [False, False, True]
    assert total["mtdPlan"] == 1775 and total["projection"] == 9318


def test_weekly_closure_targets(review):
    wk = review["targetsPage"]["weekly"]
    assert [w["label"] for w in wk["weeks"]] == ["WK-41", "WK-42", "WK-43", "WK-44"]
    assert [w["status"] for w in wk["weeks"]] == ["done", "not started", "not started", "not started"]
    assert [c["target"] for c in wk["total"]["weeks"]] == [367, 457, 431, 520]
    assert wk["total"]["frt"] == 259  # O&M FRT Count
    rj = next(x for x in wk["rows"] if x["code"] == "RJ")
    assert rj["frt"] == 208 and rj["weeks"][0]["achieved"] == 0 and rj["weeks"][1]["achieved"] is None


def test_invalid_plan_file_is_rejected_not_trusted(tmp_path, monkeypatch):
    from app.analytics import review_targets

    monkeypatch.setattr(review_targets, "get_settings", lambda: type("S", (), {"data_dir": tmp_path})())
    assert review_targets.load_plan()[0] is None  # missing file -> blank plan columns
    (tmp_path / "review_plan.json").write_text('{"monthlyPlan": {"October": {"RAJASTHAN": -5}}}', encoding="utf-8")
    plan, problem = review_targets.load_plan()
    assert plan is None and "invalid" in problem


# ---------------------------------------------------------------- Page 3: Critical KPI – Blocks & GPs


def test_critical_table_columns_follow_the_data(review):
    ct = review["sections"]["rajasthan"]["critical"]
    assert [c["label"] for c in ct["columns"]] == ["Aug-26", "Sep-26", "Sep-26 TD", "Oct-26 MTD", "WK-39", "WK-40", "WK-41"]  # newest week = reporting week
    assert ct["title"] == "Monthly & Weekly Performance – Rajasthan"


def test_critical_table_values(review, imported):
    import pandas as pd

    ct = review["sections"]["rajasthan"]["critical"]
    row = {r["key"]: r for r in ct["rows"]}
    # O&M / client-deck figures for the completed months
    assert row["gpsAbove98"]["values"][:2] == [2422, 2466]
    assert row["blocksBelow85"]["values"][:2] == [39, 37]
    assert row["gpsAtZero"]["values"][:2] == [10, 1]
    # weeks recomputed from the weekly W-nn column
    wf = imported.week_frame("W-41")
    sel = imported.gps["circle"] == "RAJASTHAN"
    v = wf["weekly"].where(wf["amc"])[sel]
    assert row["gpsAbove98"]["values"][-1] == int((v > 98).sum())
    assert row["blocksBelow85"]["values"][-1] == int((v.groupby(imported.gps.loc[sel, "block_id"]).mean() < 85).sum())
    amc = pd.to_datetime(imported.gps["amc_date"], errors="coerce")
    inc = sel & (imported.gps["amc_category"] == "INCREMENTAL") & (amc >= "2026-09-25") & (amc <= "2026-10-01")
    assert row["incrementalAdded"]["values"][-2] == int(inc.sum())  # WK-40
    assert row["gpsAbove98"]["scope"] == 4710 and row["oneTimeUpGps"]["scope"] == 8425
    assert row["oneTimeUpGps"]["source"] == "circle sheet"
    ins = {i["label"]: i for i in ct["insight"]["items"]}
    assert ins[">98% GPs"]["current"] == row["gpsAbove98"]["values"][-1]


def test_pptx_contains_every_circle_for_page3(imported, client):
    login(client)
    p = client.get("/api/v1/reviews/management.pptx", params={"week": "W-41", "compare": "W-40", "sections": "blocks-gps"})
    deck = Presentation(io.BytesIO(p.content))
    titles = [next((s.text_frame.text for s in slide.shapes if s.name == "Title"), "") for slide in deck.slides]
    assert [t for t in titles if t.startswith("Critical KPI")] == [f"Critical KPI – Blocks & GPs {c}" for c in ("Rajasthan", "Haryana", "Assam")]


def test_previous_month_td_column(review, imported, client):
    """Sep-26 TD = 1-8 Sep (same days as Oct-26 TD), same rules as the TD column; the PPT keeps its layout."""
    import pandas as pd

    from app.analytics import views
    from app.reports.review_pptx import _without_previous_td

    ptd = next(p for p in review["periods"] if p["kind"] == "ptd")
    assert ptd["label"] == "Sep-26 TD" and ptd["endDate"] == "2026-09-08"
    i = review["periods"].index(ptd)
    rj = review["sections"]["rajasthan"]
    # same figure as the Circle Insights "Sep-26 TD" column
    ci = views.circle_insights(imported, "rajasthan", "W-41", "W-40")["periodTable"]
    j = [p["kind"] for p in ci["periods"]].index("ptd")
    assert rj["gpUptime"]["bands"][0]["values"][i] == ci["gpBands"][j]["counts"][">98%"]
    # GP counted only when under AMC by 8 Sep
    amc = pd.to_datetime(imported.gps["amc_date"], errors="coerce")
    sel = imported.gps["circle"] == "RAJASTHAN"
    assert rj["gpUptime"]["scope"][i] == int((sel & (amc <= "2026-09-08")).sum())
    # Page 3: incremental GPs with AMC date 1-8 Sep; circle-sheet rows stay empty
    row = {r["key"]: r for r in rj["critical"]["rows"]}
    k = [c["kind"] for c in rj["critical"]["columns"]].index("ptd")
    inc = sel & (imported.gps["amc_category"] == "INCREMENTAL") & (amc >= "2026-09-01") & (amc <= "2026-09-08")
    assert row["incrementalAdded"]["values"][k] == int(inc.sum())
    assert row["oneTimeUpGps"]["values"][k] is None
    # the deck drops the column everywhere
    deck = _without_previous_td(review)
    assert all(p["kind"] != "ptd" for p in deck["periods"])
    assert len(deck["sections"]["rajasthan"]["gpUptime"]["total"]) == len(deck["periods"])
    assert all(c["kind"] != "ptd" for c in deck["sections"]["rajasthan"]["critical"]["columns"])


# ---------------------------------------------------------------- AMC conversion: plan vs execution page


def test_amc_conversion_plan_vs_execution(imported, plan_file, review, client):
    """Same plan and rules as Management Review page 2; months / weeks not started have no execution."""
    login(client)
    body = client.get("/api/v1/amc-conversion", params={"week": "W-41", "compare": "W-40"}).json()
    m = body["monthly"]
    assert [x["label"] for x in m["months"]] == ["Oct-26", "Nov-26", "Dec-26", "Jan-27"]
    assert m["baselineDate"] == "2026-09-30"
    assert [c["plan"] for c in m["total"]["months"]] == [1775, 2505, 2503, 1061]
    assert [c["planPartial"] for c in m["total"]["months"]] == [False, False, False, True]
    assert all(c["achieved"] is None for c in m["total"]["months"][1:])  # not started
    rm = {r["id"]: r for r in review["targetsPage"]["roadmap"]["rows"]}
    for row in m["rows"]:
        oct_ = row["months"][0]
        assert (oct_["plan"], oct_["achieved"]) == (rm[row["id"]]["mtdPlan"], rm[row["id"]]["mtdAchieved"])
        assert oct_["cumulativeAchieved"] == rm[row["id"]]["cumulativeMtd"]
        assert oct_["cumulativePlan"] == rm[row["id"]]["planByMonth"]
    assert body["headline"]["monthPlan"] == 1775 and body["headline"]["projection"] == 9318
    wk = body["weekly"]
    assert [w["label"] for w in wk["weeks"]] == ["WK-41", "WK-42", "WK-43", "WK-44"]
    assert [c["target"] for c in wk["total"]["weeks"]] == [367, 457, 431, 520]
    mr = review["targetsPage"]["weekly"]
    assert [c["achieved"] for c in wk["total"]["weeks"]][:1] == [c["achieved"] for c in mr["total"]["weeks"]][:1]


# ---------------------------------------------------------------- AMC conversion: GP-wise plan (Planned Date / Ach Date)


def test_amc_plan_summary_without_plan_columns(imported, client):
    """A file without the GP-wise plan block: KPIs are 0 with a clear warning, lists are empty – nothing breaks."""
    login(client)
    body = client.get("/api/v1/amc-plan/summary", params={"week": "W-41"}).json()
    assert body["hasPlan"] is False and body["total"]["monthTarget"] == 0 and body["warnings"]
    assert body["month"]["label"] == "Oct 2026" and body["week"]["id"] == "W-41"
    for kind in ("amc-plan-month", "amc-achieved-mtd", "amc-plan-week", "amc-achieved-week", "amc-overdue"):
        assert client.get(f"/api/v1/kpi-lists/{kind}", params={"week": "W-41"}).json()["total"] == 0


def test_amc_plan_columns_are_parsed():
    """14 Oct file: the GP-wise plan block is read as-is (counts checked against the sheet)."""
    import datetime as dt
    from pathlib import Path

    import pytest

    from app.ingestion.parser import parse_workbook

    path = Path(__file__).resolve().parents[2] / "data" / "raw" / "Master Data V8.2 Daily track 14-10-26.xlsb"
    if not path.exists():
        pytest.skip("14 Oct workbook not available")
    p = parse_workbook(path)
    plan = p.amc_plan
    assert len(plan) == 1651
    assert sum(1 for r in plan if r["planned_date"] and r["planned_date"].month == 10) == 1134
    assert sum(1 for r in plan if r["ach_date"] and dt.date(2026, 10, 1) <= r["ach_date"] <= dt.date(2026, 10, 14)) == 265
    assert {r["category"] for r in plan} == {"5% Incremental", "Additional Plan"}


def test_amc_plan_timeline_without_plan_columns(imported, client):
    """Monthly / weekly plan views on a file without the GP-wise plan: zero targets, a warning, current month and week present."""
    login(client)
    body = client.get("/api/v1/amc-plan/timeline", params={"week": "W-41"}).json()
    assert body["hasPlan"] is False and body["warnings"]
    assert [m["label"] for m in body["monthly"]["months"]] == ["Oct 2026"] and body["monthly"]["months"][0]["current"]
    assert any(w["current"] and w["id"] == "W-41" for w in body["weekly"]["weeks"])
    assert all(c["target"] == 0 for c in body["weekly"]["total"]["cells"])


def test_amc_plan_summary_picked_date(imported, client):
    """A picked date gives that day's month and Fri–Thu week; the default equals picking the data date of the week."""
    login(client)
    default = client.get("/api/v1/amc-plan/summary", params={"week": "W-41"}).json()
    assert default["dates"][0] == imported.data_as_of.isoformat()
    picked = client.get("/api/v1/amc-plan/summary", params={"week": "W-41", "date": default["asOf"]}).json()
    assert picked["total"] == default["total"] and picked["week"] == default["week"]
    sept = client.get("/api/v1/amc-plan/summary", params={"date": "2026-09-30"}).json()
    assert sept["asOf"] == "2026-09-30" and sept["month"]["label"] == "Sep 2026" and sept["week"]["id"] == "W-40"
    assert client.get("/api/v1/kpi-lists/amc-plan-week", params={"date": "2026-09-30"}).json()["total"] == 0
    assert client.get("/api/v1/amc-plan/summary", params={"date": "not-a-date"}).status_code == 422


def test_month_targets_switch_from_plan_file_to_master_data(imported, monkeypatch):
    """A state-month without Planned Dates uses the targets file ("plan"); once the file has Planned Dates for it, the
    target is their count ("excel") – no other state or month changes."""
    import dataclasses
    import datetime as dt

    import pandas as pd

    from app.analytics import amc_plan

    targets = amc_plan._MonthTargets(source="test", targets={"2026-10": {"RAJASTHAN": 1063, "HARYANA": 645}, "2026-11": {"RAJASTHAN": 1063, "HARYANA": 1242}})
    monkeypatch.setattr(amc_plan, "load_month_targets", lambda: (targets, None))
    hr = imported.gps.index[imported.gps["circle"] == "HARYANA"][:3]
    plan = pd.DataFrame(
        {
            "category": ["Additional Plan"] * 3,
            "planned_date": [dt.date(2026, 11, 2), dt.date(2026, 11, 3), dt.date(2026, 10, 5)],
            "ach_date": [None, None, dt.date(2026, 10, 6)],
            "current_status": [None] * 3,
            "letter_status": [None] * 3,
            "letter_date": [None] * 3,
        },
        index=hr,
    )
    st = dataclasses.replace(imported, amc_plan=plan, _memo={}, _memo_locks={})
    r = amc_plan.amc_plan_timeline(st, "W-41")
    months = [m["id"] for m in r["monthly"]["months"]]
    cell = lambda name, month: next(x for x in r["monthly"]["rows"] if x["name"] == name)["cells"][months.index(month)]  # noqa: E731
    assert (cell("Haryana", "2026-10")["target"], cell("Haryana", "2026-10")["targetSource"]) == (1, "excel")  # file now has 1 Oct Planned Date
    assert (cell("Haryana", "2026-11")["target"], cell("Haryana", "2026-11")["targetSource"]) == (2, "excel")
    assert cell("Haryana", "2026-10")["achieved"] == 1
    assert (cell("Rajasthan", "2026-11")["target"], cell("Rajasthan", "2026-11")["targetSource"]) == (1063, "plan")  # untouched
    total_nov = r["monthly"]["total"]["cells"][months.index("2026-11")]
    assert total_nov["target"] == 1063 + 2 and total_nov["targetSource"] == "mixed"
